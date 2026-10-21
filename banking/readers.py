"""Turning an uploaded statement file into rows of text cells (D-064, widened in D-211).

A statement arrives as CSV (comma, semicolon or tab), Excel (.xlsx) or a text PDF, whichever the bank or
the M-Pesa portal gives. `tables` returns the candidate tables found in the file, each a list of rows; the
caller picks the one with a recognisable header row. Pure functions over the file's bytes; no database.
"""

import csv
import datetime
import io
import re
import statistics
import zipfile
from decimal import Decimal

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PyPdfError

from core import uploads

DELIMITERS = (",", ";", "\t")
MAX_PDF_PAGES = 100
MAX_SHEETS = 5
MAX_COLUMNS = 40

# A date at the start of a cell: 14-Oct-26, 14/10/2026, 14.10.26, 14 October 2026, 2026-10-14.
_DATE_START = re.compile(r"^\s*(\d{1,2}[-/. ](?:[A-Za-z]{3,9}|\d{1,2})[-/. ,]*\d{2,4}|\d{4}-\d{2}-\d{2})\b")


def starts_with_date(cell: str) -> bool:
    return bool(_DATE_START.match(cell or ""))


class FileError(ValueError):
    """The file cannot be opened as a statement at all."""


def tables(data: bytes, names: dict[str, tuple[str, ...]], anchors: tuple[str, ...]) -> list[list[list[str]]]:
    """The candidate tables in the file. `names` are the header names the caller knows, normalised by
    `normalise`; `anchors` are the fields whose column holds a date, so a PDF row can be told from its
    wrapped lines."""
    if data.startswith(b"%PDF"):
        return [_pdf_table(data, names, anchors)]
    if data.startswith(b"PK\x03\x04"):
        return _xlsx_tables(data)
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        raise FileError("Old .xls files are not supported. In Excel choose Save As, then .xlsx or CSV.")
    text = _decode(data)
    if "<table" in text[:5000].lower() or text.lstrip()[:15].lower().startswith(("<!doctype", "<html")):
        raise FileError("This is a web page, not a statement. Download the statement as CSV, Excel or PDF.")
    return [list(csv.reader(io.StringIO(text), delimiter=d)) for d in DELIMITERS]


def normalise(value: str) -> str:
    words = re.sub(r"[^a-z0-9]+", " ", (value or "").lower()).split()
    if words and words[-1] in ("kes", "ksh", "kshs"):
        words = words[:-1]
    return " ".join(words)


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        if encoding == "utf-16" and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
            continue
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def _cell_text(value) -> str:
    """An Excel cell as the text a CSV would hold: dates as 2026-01-31, numbers without float noise."""
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S") if value.time() != datetime.time() else value.strftime("%Y-%m-%d")
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, float):
        return format(Decimal(repr(value)), "f")
    return str(value).strip()


def _xlsx_tables(data: bytes) -> list[list[list[str]]]:
    if uploads.zip_is_unsafe(io.BytesIO(data)):
        raise FileError("This does not look like an Excel (.xlsx) file. Save it as .xlsx or as CSV.")
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return [[[_cell_text(v) for v in record] for record in sheet.iter_rows(values_only=True, max_col=MAX_COLUMNS)]
                for sheet in book.worksheets[:MAX_SHEETS]]
    except (zipfile.BadZipFile, InvalidFileException, KeyError, IndexError, OSError, ValueError):
        raise FileError("This does not look like an Excel (.xlsx) file. Save it as .xlsx or as CSV.") from None


# ---------------------------------------------------------------------------
# PDF
#
# A bank app or portal prints the statement as a drawn table. The text of each cell comes out with its
# position, so the table is rebuilt from where things sit on the page: the header row gives the columns,
# a line with a date in the date column starts a transaction, and the other lines (a description that
# wraps, or sits above and below a vertically centred row) belong to the nearest transaction.
# ---------------------------------------------------------------------------

_LINE_TOLERANCE = 2.5


def _chunks(page) -> list[tuple[float, float, float, str]]:
    """(x, y, width, text) for each piece of text on the page, y growing upward."""
    found = []

    def visit(text, cm, tm, _font, size):
        text = " ".join((text or "").split())
        if not text:
            return
        x = tm[4] * cm[0] + tm[5] * cm[2] + cm[4]
        y = tm[4] * cm[1] + tm[5] * cm[3] + cm[5]
        scale = abs(tm[0] * cm[0]) or 1
        found.append((x, y, len(text) * (size or 9) * 0.5 * (scale if scale < 3 else 1), text))

    page.extract_text(visitor_text=visit)
    return found


def _lines(chunks):
    """Chunks grouped into visual lines, top of the page first; each line sorted left to right."""
    lines = []
    for chunk in sorted(chunks, key=lambda c: -c[1]):
        if lines and abs(lines[-1][0] - chunk[1]) <= _LINE_TOLERANCE:
            lines[-1][1].append(chunk)
        else:
            lines.append([chunk[1], [chunk]])
    return [(y, sorted(items)) for y, items in lines]


def _header_cells(line, names) -> list[tuple[float, str, str]] | None:
    """(centre, text, normalised) of the header cells when this line is a header row."""
    cells = []
    for x, _y, width, text in line:
        if cells and x - cells[-1][3] < 6:  # a header split in two pieces: "Money" "In"
            previous = cells[-1]
            cells[-1] = (previous[0], previous[1] + " " + text, x + width, x + width)
        else:
            cells.append((x, text, x + width, x + width))
    known = {option for options in names.values() for option in options}
    out = [((start + end) / 2, text, normalise(text)) for start, text, end, _ in cells]
    return out if sum(1 for _c, _t, n in out if n in known) >= 3 else None


def _pdf_table(data: bytes, names, anchors) -> list[list[str]]:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise FileError("This PDF is password protected. Open it, remove the password and save it again.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise FileError(f"The PDF has more than {MAX_PDF_PAGES} pages; download it by month.")
        pages = [_lines(_chunks(page)) for page in reader.pages]
    except FileError:
        raise
    except (PyPdfError, ValueError, KeyError, IndexError, TypeError, OSError, RecursionError):
        raise FileError("This PDF could not be read. Download the statement again, or as Excel or CSV.") from None
    if not any(pages):
        raise FileError("This PDF has no text to read (it looks like a scan or a screenshot). Download the "
                        "statement from your bank's app or portal as Excel or CSV, or as a PDF.")
    header, rows = None, []
    for lines in pages:
        header = next((h for h in (_header_cells(items, names) for _y, items in lines) if h), header)
        if header is None:
            continue
        rows.extend(_page_rows(lines, header, names, anchors))
    if header is None:
        return []
    return [[text for _c, text, _n in header], *rows]


def _page_rows(lines, header, names, anchors) -> list[list[str]]:
    centres = [c for c, _t, _n in header]
    date_columns = {i for i, (_c, _t, n) in enumerate(header)
                    if any(n in names.get(key, ()) for key in anchors)}

    def column(x, width):
        centre = x + width / 2
        return min(range(len(centres)), key=lambda i: abs(centres[i] - centre))

    placed = []  # (y, column, x, text)
    for y, items in lines:
        if _header_cells(items, names):
            continue
        for x, _y, width, text in items:
            placed.append((y, column(x, width), x, text))
    anchor_ys = sorted({y for y, col, _x, text in placed if col in date_columns and _DATE_START.match(text)},
                       reverse=True)
    if not anchor_ys:
        return []
    gaps = [a - b for a, b in zip(anchor_ys, anchor_ys[1:], strict=False)]
    reach = (statistics.median(gaps) / 2 if gaps else 14) + 1
    rows = {y: [[] for _ in centres] for y in anchor_ys}
    for y, col, _x, text in sorted(placed, key=lambda p: (-p[0], p[2])):
        nearest = min(anchor_ys, key=lambda a: abs(a - y))
        if abs(nearest - y) <= reach:
            rows[nearest][col].append(text)
    return [[" ".join(cell) for cell in rows[y]] for y in anchor_ys]
