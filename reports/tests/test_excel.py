"""Excel downloads on the finance pages (D-132): the CSV rows in a workbook, same rights, text never a formula."""

import io

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from accounts.tests.factories import add_member, make_org, make_property
from accounts.tests.test_isolation import login
from reports import xlsx
from reports.tests.test_income import make_lease

pytestmark = pytest.mark.django_db


def cells(response):
    sheet = load_workbook(io.BytesIO(response.content)).active
    return [c.value for row in sheet.iter_rows() for c in row]


def test_numbers_become_numbers_and_codes_stay_text():
    from django.http import HttpResponse

    r = HttpResponse("﻿Name,Amount,Phone\n=SUM(A1),1500.50,0712345678\n", content_type="text/csv")
    r["Content-Disposition"] = 'attachment; filename="t.csv"'
    out = xlsx.from_csv(r)
    sheet = load_workbook(io.BytesIO(out.content)).active
    assert sheet["A2"].value == "=SUM(A1)" and sheet["A2"].data_type == "s"
    assert sheet["B2"].value == 1500.5 and sheet["C2"].value == "0712345678"


@pytest.mark.parametrize("name", ["profit_loss", "cash_flow", "fees", "units", "receivables", "dashboard_export"])
def test_finance_pages_offer_excel(client, name):
    owner = make_org(name="Acme")
    make_lease(owner, make_property(owner.organization, name="Acacia"), code="A1")
    login(client, owner)
    response = client.get(reverse(f"reports:{name}"), {"format": "xlsx"})
    assert response.status_code == 200 and response["Content-Type"] == xlsx.CONTENT_TYPE
    assert response["Content-Disposition"].endswith('.xlsx"')


def test_excel_needs_the_export_right(client):
    owner = make_org(name="Acme")
    prop = make_property(owner.organization, name="Acacia")
    viewer = add_member(owner.organization, "caretaker", properties=[prop])
    login(client, viewer)
    assert client.get(reverse("reports:cash_flow"), {"format": "xlsx"}).status_code == 403
