"""The payments received report (D-132): states, methods, the waiting queue, isolation and rights."""

import datetime

import pytest
from django.core.exceptions import PermissionDenied
from django.urls import reverse
from django.utils import timezone

from accounts.tests.factories import add_member, make_org, make_property
from accounts.tests.test_isolation import login
from payments.models import Payment
from reports import operations
from reports.tests.test_income import make_lease

pytestmark = pytest.mark.django_db

D = datetime.date


@pytest.fixture
def owner():
    return make_org(name="Acme")


@pytest.fixture
def prop(owner):
    return make_property(owner.organization, name="Acacia Court")


@pytest.fixture
def lease(owner, prop):
    return make_lease(owner, prop, code="A1")


def pay(owner, lease, amount, day, status, confirmed=True, method="MPESA"):
    p = Payment.objects.create(organization=owner.organization, lease=lease, amount=amount, method=method,
                               paid_at=day, recorded_by=owner.user)
    stamp = timezone.make_aware(datetime.datetime.combine(day, datetime.time(9))) if confirmed else None
    Payment.objects.filter(pk=p.pk).update(status=status, confirmed_at=stamp)
    return p


def test_states_and_only_confirmed_is_received(owner, lease):
    pay(owner, lease, 1000, D(2026, 3, 5), "CONFIRMED")
    pay(owner, lease, 2000, D(2026, 3, 6), "CONFIRMED", method="CASH")
    pay(owner, lease, 400, D(2026, 3, 7), "PENDING_REVIEW", confirmed=False)
    pay(owner, lease, 300, D(2026, 3, 8), "REVERSED")  # confirmed, then undone
    pay(owner, lease, 200, D(2026, 3, 9), "REVERSED", confirmed=False)  # turned down in review
    pay(owner, lease, 999, D(2026, 2, 9), "CONFIRMED")  # another month
    r = operations.payments(owner, D(2026, 3, 1), D(2026, 3, 31))
    by = {k: (n, a) for k, _label, n, a in r.states}
    assert by == {"pending": (1, 400), "confirmed": (2, 3000), "reversed": (1, 300), "rejected": (1, 200)}
    assert r.confirmed == 3000 and r.undone == 500
    assert [(label, n) for label, n, _a, _s in r.methods] == [("Cash", 1), ("M-Pesa", 1)]
    assert sum(s for *_x, s in r.methods) == 1
    assert (r.waiting, r.waiting_amount, r.oldest_waiting) == (1, 400, D(2026, 3, 7))


def test_empty_and_other_organizations_are_left_out(owner, lease):
    other = make_org(name="Other")
    pay(other, make_lease(other, make_property(other.organization), code="Z1"), 5000, D(2026, 3, 5), "CONFIRMED")
    r = operations.payments(owner, D(2026, 3, 1), D(2026, 3, 31))
    assert r.total == 0 and r.confirmed == 0 and r.methods == [] and r.waiting == 0


def test_a_property_filter_narrows_the_figures(owner, prop, lease):
    second = make_property(owner.organization, name="Other Court")
    pay(owner, lease, 1000, D(2026, 3, 5), "CONFIRMED")
    pay(owner, make_lease(owner, second, code="B1"), 700, D(2026, 3, 5), "CONFIRMED")
    assert operations.payments(owner, D(2026, 3, 1), D(2026, 3, 31)).confirmed == 1700
    assert operations.payments(owner, D(2026, 3, 1), D(2026, 3, 31), prop).confirmed == 1000


def test_needs_the_right(owner, prop):
    caretaker = add_member(owner.organization, "caretaker", properties=[prop])
    with pytest.raises(PermissionDenied):
        operations.payments(caretaker, D(2026, 3, 1), D(2026, 3, 31))


def test_page_and_csv_match(client, owner, lease):
    pay(owner, lease, 1000, timezone.localdate(), "CONFIRMED")
    login(client, owner)
    page = client.get(reverse("reports:payments"))
    assert page.status_code == 200 and "status=confirmed" in page.content.decode()
    body = client.get(reverse("reports:payments"), {"format": "csv"}).content.decode("utf-8-sig")
    assert "Confirmed,1,1000" in body.replace(" ", "")


def test_excel_export_is_audited_and_keeps_text_as_text(client, owner, lease):
    import io

    from openpyxl import load_workbook

    from audit.models import AuditEvent

    pay(owner, lease, 1000, timezone.localdate(), "CONFIRMED")
    login(client, owner)
    response = client.get(reverse("reports:payments"), {"format": "xlsx"})
    assert response.status_code == 200 and response["Content-Disposition"].endswith('.xlsx"')
    sheet = load_workbook(io.BytesIO(response.content)).active
    values = [c.value for row in sheet.iter_rows() for c in row]
    assert "Confirmed" in values and 1000 in values
    assert AuditEvent.objects.filter(action="reports.export", changes__format__1="xlsx").exists()


def test_month_bars_link_to_the_payments_list_with_the_filters(client, owner, lease):
    pay(owner, lease, 1000, timezone.localdate(), "CONFIRMED")
    login(client, owner)
    html = client.get(reverse("reports:payments"), {"property": str(lease.unit.property.public_id)}).content.decode()
    assert "status=confirmed" in html and "date_from=" in html and "property=" in html
