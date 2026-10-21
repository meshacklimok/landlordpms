"""Excel downloads (D-132). Every report that offers CSV also offers Excel: the same rows, one workbook sheet."""

import csv
import datetime
import io
import re
from decimal import Decimal
from functools import wraps

from django.http import HttpResponse

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# Plain numbers become numbers; anything with a leading zero (a phone number, a code) stays text.
_NUMBER = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?$")


_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DATETIME = re.compile(r"^\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2})?$")
PRIMARY = "146C4A"
TINT = "EEF4F0"
LINE = "D9DFDB"
MAX_WIDTH = 48
MIN_WIDTH = 9


def _value(v):
    """A cell value Excel can hold: money as float, ISO dates as real dates, everything else as given."""
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, str):
        try:
            if _DATE.match(v):
                return datetime.date.fromisoformat(v)
            if _DATETIME.match(v):
                return datetime.datetime.fromisoformat(v.replace("T", " "))
        except ValueError:
            return v
    return v


def _is_text(v) -> bool:
    return isinstance(v, str) and bool(v.strip()) and not _NUMBER.match(v) and not _DATE.match(v)


def _header_row(rows) -> int | None:
    """Index of the row that heads the table: the first of the top rows that is as wide as the widest and all words."""
    width = max((len(r) for r in rows), default=0)
    if width < 2:
        return None
    for i, row in enumerate(rows[:8]):
        cells = [c for c in row if c not in ("", None)]
        if (len(row) == width and cells and all(_is_text(c) for c in cells) and (width > 2 or i == 0)
                and len(cells) >= 2):
            return i
    return None


def _width(values) -> float:
    def size(v):
        if isinstance(v, float):
            return len(f"{v:,.2f}")
        return max((len(part) for part in str(v).split("\n")), default=0)

    longest = max((size(v) for v in values if v not in ("", None)), default=0)
    return min(max(longest + 3, MIN_WIDTH), MAX_WIDTH)


def response(filename: str, title: str, rows: list[list]) -> HttpResponse:
    """The rows as a workbook, laid out to read and print: styled header, widths that fit, money and dates formatted,
    totals in bold, header frozen and repeated on every printed page. Text is always stored as text, so a cell that
    starts with = is never a formula."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import TYPE_STRING
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = "".join(" " if ch in "[]*?:/" + chr(92) else ch for ch in title)[:31] or "Report"
    book.properties.creator = "NyumbaPMS"
    book.properties.title = sheet.title
    rows = [list(r) for r in rows]
    head = _header_row(rows)
    for row in rows:
        sheet.append([_value(v) for v in row])
    columns = max((len(r) for r in rows), default=0)
    thin = Side(style="thin", color=LINE)
    for r, row in enumerate(rows, start=1):
        label = str(row[0]).strip().lower() if row and row[0] not in (None, "") else ""
        is_head, is_total = head is not None and r == head + 1, label.startswith(("total", "net ", "grand"))
        for c, cell in enumerate(sheet[r], start=1):
            if isinstance(cell.value, str):
                cell.data_type = TYPE_STRING
            elif isinstance(cell.value, datetime.datetime):
                cell.number_format = "dd mmm yyyy hh:mm"
            elif isinstance(cell.value, datetime.date):
                cell.number_format = "dd mmm yyyy"
            elif isinstance(cell.value, float) and cell.value != int(cell.value):
                cell.number_format = "#,##0.00"
            elif isinstance(cell.value, int | float) and abs(cell.value) >= 10000:
                cell.number_format = "#,##0"
            if is_head:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=PRIMARY)
                numeric = sum(isinstance(x[c - 1], int | float | Decimal) for x in rows[head + 1:] if len(x) >= c)
                cell.alignment = Alignment(wrap_text=True, vertical="center",
                                           horizontal="right" if numeric > len(rows[head + 1:]) / 2 else "left")
            elif is_total:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor=TINT)
                cell.border = Border(top=Side(style="thin", color=PRIMARY))
            elif head is not None and r > head + 1:
                cell.border = Border(bottom=thin)
                if isinstance(cell.value, str) and "\n" in cell.value:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            elif r <= (head if head is not None else 0) and c == 1:
                cell.font = Font(bold=True, color=PRIMARY)
    for c in range(1, columns + 1):
        letter = get_column_letter(c)
        values = [x[c - 1] for x in rows if len(x) >= c]
        # Preamble rows (scope, period) sit above the table and must not stretch its first column.
        if head is not None:
            values = [x[c - 1] for x in rows[head:] if len(x) >= c]
        sheet.column_dimensions[letter].width = _width(values)
    if head is not None:
        sheet.row_dimensions[head + 1].height = 30
        sheet.freeze_panes = sheet.cell(row=head + 2, column=1)
        if len(rows) > head + 2:
            sheet.auto_filter.ref = f"A{head + 1}:{get_column_letter(columns)}{len(rows)}"
        sheet.print_title_rows = f"{head + 1}:{head + 1}"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape" if columns > 7 else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = sheet.page_margins.right = 0.4
    sheet.oddHeader.left.text, sheet.oddHeader.right.text = sheet.title, "&D"
    sheet.oddFooter.left.text, sheet.oddFooter.right.text = "NyumbaPMS", "Page &P of &N"
    out = HttpResponse(content_type=CONTENT_TYPE)
    out["Content-Disposition"] = f'attachment; filename="{filename}"'
    book.save(out)
    return out


def from_csv(csv_response: HttpResponse) -> HttpResponse:
    text = csv_response.content.decode("utf-8-sig")
    # The CSV quotes text that looks like a formula with a leading apostrophe; Excel text cells need no such guard.
    rows = [[float(c) if _NUMBER.match(c) else (c[1:] if c[:2] in ("'=", "'+", "'-", "'@") else c) for c in row]
            for row in csv.reader(io.StringIO(text))]
    name = re.search(r'filename="([^"]+)"', csv_response.get("Content-Disposition", ""))
    stem = re.sub(r"\.csv$", "", name.group(1)) if name else "report"
    return response(f"{stem}.xlsx", stem, rows)


def also_excel(view):
    """Wrap a view so `?format=xlsx` runs it as `?format=csv` (same rights, same audit) and converts the result."""

    @wraps(view)
    def wrapped(request, *args, **kwargs):
        if request.GET.get("format") != "xlsx":
            return view(request, *args, **kwargs)
        query = request.GET.copy()
        query["format"] = "csv"
        request.GET = query
        result = view(request, *args, **kwargs)
        if result.status_code == 200 and result.get("Content-Type", "").startswith("text/csv"):
            return from_csv(result)
        return result

    for attr in ("view_class", "view_initkwargs"):
        if hasattr(view, attr):
            setattr(wrapped, attr, getattr(view, attr))
    return wrapped
