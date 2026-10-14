"""Bulk import of units, tenants, tenants moving in (tenant + lease) and opening balances (doc 11 §20, doc 14 A2/B8,
D-156). The rows come from a CSV or Excel file, or from a photo of the printed form filled in by hand.

Flow: upload -> preview -> apply -> (undo within 24 hours).
- The preview runs every row through the normal create services inside a transaction that is
  rolled back, so it reports exactly what applying would do, duplicates within the file included.
- Applying runs the rows again for real. Rows that fail are skipped and reported; the rest are kept.
- Undo deletes the units and tenants the batch created, except any already used on a lease. For opening
  balances it reverses each entry that is still the lease's balance brought forward; the ledger keeps both.
- A photo is read by the model (imports/scan.py) into the same rows, so it gets the same preview; the person must
  tick that they checked the rows against the paper before it is applied.
- Move-in rows create the tenant (or reuse the one with that phone) and a lease, left as a draft unless the batch
  was uploaded with "activate". Undo removes drafts and the tenants made for them; an issued lease is kept.
- Codes and phones are normalised by the same services the forms use.
- A batch belongs to whoever uploaded it. Members with every property can also act on anyone's batch;
  members limited to some properties see only their own.
- A preview holds the raw file, ID numbers included, so it expires after PREVIEW_TTL and
  ``purge_stale_previews`` (run daily) empties it.
"""

import csv
import datetime
import io
import re
import zipfile

from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.db.models import ProtectedError
from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from accounts.models import Membership
from accounts.permissions import can, require, visible_properties
from audit import services as audit
from billing import invoicing
from billing.models import LedgerEntry
from core.money import parse_money
from core.phone import InvalidPhoneNumber, normalize_phone
from leases import services as lease_services
from leases.models import Lease
from properties import services as property_services
from properties.models import Building, Property, Unit, clean_code
from tenants import services as tenant_services
from tenants.models import Tenant

from .models import ImportBatch

Kind = ImportBatch.Kind
Status = ImportBatch.Status

MAX_BYTES = 1_000_000
MAX_XLSX_BYTES = 3_000_000
MAX_ROWS = 2000
PREVIEW_TTL = datetime.timedelta(hours=24)

COLUMNS = {
    Kind.UNITS: ("property_code", "unit_code", "building", "unit_type", "type_label", "list_rent"),
    Kind.TENANTS: ("name", "phone", "alt_phone", "email", "kind", "contact_person", "id_type", "id_number",
                   "kra_pin", "emergency_contact_name", "emergency_contact_phone", "notes"),
    Kind.BALANCES: ("property_code", "unit_code", "lease_number", "amount", "as_of", "reason"),
    Kind.LEASES: ("property_code", "unit_code", "name", "phone", "email", "rent", "deposit", "start_date", "due_day"),
}
REQUIRED = {
    Kind.UNITS: ("property_code", "unit_code"), Kind.TENANTS: ("name", "phone"),
    Kind.BALANCES: ("property_code", "unit_code", "amount", "as_of"),
    Kind.LEASES: ("property_code", "unit_code", "name", "phone", "rent", "start_date"),
}
# A row where all of these are empty is a unit or lease the person left alone on a sheet we pre-filled.
SKIP_IF_EMPTY = {Kind.LEASES: ("name", "phone"), Kind.BALANCES: ("amount",)}
PHONE_COLUMNS = ("phone", "alt_phone", "emergency_contact_phone")
EXAMPLES = {
    Kind.UNITS: ("GV", "A1", "Block A", "APARTMENT", "2 bedroom", "15000"),
    Kind.TENANTS: ("Wanjiku Kamau", "0712345678", "", "wanjiku@example.com", "INDIVIDUAL", "", "NATIONAL_ID",
                   "12345678", "A123456789B", "", "", ""),
    Kind.BALANCES: ("GV", "A1", "", "12500", "2026-01-31", "Arrears from the old spreadsheet"),
    Kind.LEASES: ("GV", "A1", "Wanjiku Kamau", "0712345678", "wanjiku@example.com", "15000", "15000", "2026-01-01",
                  "5"),
}
ALIASES = {
    "property": "property_code", "unit": "unit_code", "code": "unit_code", "type": "unit_type",
    "label": "type_label", "rent": "list_rent", "asking_rent": "list_rent",
    "tenant": "name", "tenant_name": "name", "full_name": "name", "phone_number": "phone", "mobile": "phone",
    "kra": "kra_pin", "id": "id_number", "id_no": "id_number",
    "lease": "lease_number", "balance": "amount", "opening_balance": "amount", "date": "as_of",
    "monthly_rent": "rent", "deposit_amount": "deposit", "move_in": "start_date", "move_in_date": "start_date",
    "moved_in": "start_date", "start": "start_date", "tenant_phone": "phone", "tenant_email": "email",
    "rent_due_day": "due_day", "due": "due_day",
}
CAPABILITY = {Kind.UNITS: "units.manage", Kind.TENANTS: "tenants.manage", Kind.BALANCES: "invoices.adjust",
              Kind.LEASES: "leases.draft"}
DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y")
SENSITIVE = set(Tenant.SENSITIVE_FIELDS)


def can_import(membership: Membership, kind: str) -> bool:
    return can(membership, CAPABILITY[kind])


def visible_batches(membership: Membership, queryset=None):
    """Batches this membership may see: its own, or everyone's with access to every property."""
    qs = (queryset if queryset is not None else ImportBatch.objects.all()).filter(
        organization=membership.organization, kind__in=[k for k in Kind.values if can_import(membership, k)])
    return qs if membership.all_properties else qs.filter(created_by_id=membership.user_id)


def can_access(membership: Membership, batch: ImportBatch) -> bool:
    return (batch.organization_id == membership.organization_id and can_import(membership, batch.kind)
            and (membership.all_properties or batch.created_by_id == membership.user_id))


def is_expired(batch: ImportBatch) -> bool:
    return batch.status == Status.PREVIEW and timezone.now() >= batch.created_at + PREVIEW_TTL


def purge_stale_previews() -> int:
    """Discards previews older than PREVIEW_TTL and drops their rows. Returns how many."""
    return ImportBatch.objects.filter(status=Status.PREVIEW, created_at__lt=timezone.now() - PREVIEW_TTL).update(
        status=Status.DISCARDED, rows=[], updated_at=timezone.now())


def template_csv(kind: str) -> str:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(COLUMNS[kind])
    writer.writerow(EXAMPLES[kind])
    return out.getvalue()


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------


def _header(name, kind: str) -> str:
    key = re.sub(r"[\s\-/]+", "_", str(name or "").strip().lower())
    return key if key in COLUMNS[kind] else ALIASES.get(key, key)


def _cell_text(value) -> str:
    """An Excel cell as the text a CSV would hold: dates as 2026-01-31, whole numbers without ".0"."""
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _fix_phone(value: str) -> str:
    """Excel drops the leading 0 of a number typed into a General cell: 712345678 becomes 0712345678."""
    return "0" + value if value.isdigit() and len(value) == 9 and value[0] in "71" else value


def _read_csv(upload):
    raw = upload.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise ValidationError(_("The file is too big. Split it into files under 1 MB."))
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")  # Excel's "CSV" on Windows
    reader = csv.reader(io.StringIO(text))
    try:
        for record in reader:
            yield reader.line_num, record
    except csv.Error:
        raise ValidationError(_("This does not look like a CSV file.")) from None


def _read_xlsx(upload):
    raw = upload.read(MAX_XLSX_BYTES + 1)
    if len(raw) > MAX_XLSX_BYTES:
        raise ValidationError(_("The file is too big. Split it into files under 3 MB."))
    try:
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = book.worksheets[0]
        for number, record in enumerate(sheet.iter_rows(values_only=True, max_col=30), start=1):
            if number > MAX_ROWS + 50:
                raise ValidationError(_("Import at most %(n)s rows at a time.") % {"n": MAX_ROWS})
            yield number, [_cell_text(v) for v in record]
    except (zipfile.BadZipFile, InvalidFileException, KeyError, IndexError, OSError, ValueError):
        raise ValidationError(_("This does not look like an Excel (.xlsx) file. Save it as .xlsx or as CSV.")) from None


def parse_upload(kind: str, upload) -> list[dict]:
    """Rows as [{"line", "data", "errors"}] from a .csv or .xlsx upload. Problems with the whole file raise
    ValidationError."""
    name = getattr(upload, "name", "").lower()
    if name.endswith(".xls"):
        raise ValidationError(_("Old .xls files are not supported. In Excel choose Save As, then .xlsx or CSV."))
    records = _read_xlsx(upload) if name.endswith((".xlsx", ".xlsm")) else _read_csv(upload)
    header, names, rows = None, [], []
    for line, record in records:
        if not any(str(v).strip() for v in record):
            continue
        if header is None:
            header = record
            names = [_header(h, kind) for h in header]
            unknown = [str(h) for h, n in zip(header, names, strict=True) if n and n not in COLUMNS[kind]]
            if unknown:
                raise ValidationError(_("Unknown columns: %(cols)s. Use the template's column names.")
                                      % {"cols": ", ".join(unknown)})
            missing = [c for c in REQUIRED[kind] if c not in names]
            if missing:
                raise ValidationError(_("Missing columns: %(cols)s.") % {"cols": ", ".join(missing)})
            continue
        row = _row(kind, line, {n: str(v).strip() for n, v in zip(names, record, strict=False) if n})
        if row is None:
            continue
        if len(rows) == MAX_ROWS:
            raise ValidationError(_("Import at most %(n)s rows at a time.") % {"n": MAX_ROWS})
        rows.append(row)
    if header is None:
        raise ValidationError(_("The file is empty."))
    if not rows:
        raise ValidationError(_("The file has no rows to import."))
    return rows


def make_row(kind: str, line: int, data: dict) -> dict | None:
    """One row ready to check: its required columns tested, or None for a row left empty on a pre-filled sheet."""
    skip = SKIP_IF_EMPTY.get(kind)
    if skip and not any(data.get(c) for c in skip):
        return None
    for c in PHONE_COLUMNS:
        if data.get(c):
            data[c] = _fix_phone(data[c])
    errors = [_("%(col)s is required.") % {"col": c} for c in REQUIRED[kind] if not data.get(c)]
    return {"line": line, "data": data, "errors": errors}


_row = make_row


def _choice(choices, value: str, default=""):
    """Matches a choice by value or label, ignoring case, spaces and hyphens."""
    if not value:
        return default
    key = re.sub(r"[\s\-]+", "_", value.strip()).upper()
    for v, label in choices.choices:
        if key in (v, re.sub(r"[\s\-]+", "_", str(label)).upper()):
            return v
    raise ValidationError(_("“%(value)s” is not a valid choice.") % {"value": value})


def _money(value: str):
    if not value:
        return None
    try:
        return parse_money(value)
    except ValidationError as e:
        raise ValidationError(f"“{value}”: {e.messages[0]}") from None


def _date(value: str) -> datetime.date:
    for fmt in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    raise ValidationError(_("“%(value)s” is not a date. Use 2026-01-31 or 31/01/2026.") % {"value": value})


# ---------------------------------------------------------------------------
# One row
# ---------------------------------------------------------------------------


def _create_unit(actor: Membership, data: dict, request, ctx):
    code = clean_code(data["property_code"])
    prop = visible_properties(actor, Property.objects.all()).filter(code__iexact=code).first()
    if prop is None:
        raise ValidationError(_("No property with code %(code)s.") % {"code": code})
    if not can(actor, "units.manage", prop):
        raise ValidationError(_("You cannot add units to %(name)s.") % {"name": prop.name})
    building = None
    if data.get("building"):
        building = Building.objects.filter(property=prop, name__iexact=data["building"]).first()
        if building is None:
            raise ValidationError(_("%(name)s has no building called “%(b)s”. Add it first.")
                                  % {"name": prop.name, "b": data["building"]})
    unit = property_services.create_unit(
        actor, prop, code=data["unit_code"], building=building, request=request,
        unit_type=_choice(Unit.Type, data.get("unit_type", ""), Unit.Type.APARTMENT),
        type_label=data.get("type_label", ""), list_rent=_money(data.get("list_rent", "")),
    )
    return unit, unit.payment_reference


def _create_tenant(actor: Membership, data: dict, request, ctx):
    if any(data.get(f) for f in SENSITIVE) and not can(actor, "tenants.view_sensitive"):
        raise ValidationError(_("You may not import ID numbers or KRA PINs. Leave those columns empty."))
    try:
        phone = normalize_phone(data["phone"])
    except InvalidPhoneNumber as exc:
        raise ValidationError(str(exc)) from None
    if tenant_services.tenants_with_phone(actor.organization, phone).exists():
        raise ValidationError(_("A tenant with phone %(phone)s already exists.") % {"phone": phone})
    fields = {k: data.get(k, "") for k in COLUMNS[Kind.TENANTS] if k not in ("name", "phone", "kind", "id_type")}
    tenant = tenant_services.create_tenant(
        actor, name=data["name"], phone=phone, request=request,
        kind=_choice(Tenant.Kind, data.get("kind", ""), Tenant.Kind.INDIVIDUAL),
        id_type=_choice(Tenant.IdType, data.get("id_type", "")), **fields,
    )
    return tenant, tenant.name


def _unit(actor: Membership, data: dict) -> Unit:
    code = clean_code(data["property_code"])
    unit = (Unit.objects.filter(property__in=visible_properties(actor, Property.objects.all()),
                                property__code__iexact=code, code__iexact=clean_code(data["unit_code"]))
            .select_related("property").first())
    if unit is None:
        raise ValidationError(_("No unit %(unit)s in property %(code)s.") % {"unit": data["unit_code"], "code": code})
    return unit


def _create_balance(actor: Membership, data: dict, request, ctx):
    """A balance brought forward for the unit's active lease, or for the lease named in lease_number."""
    unit = _unit(actor, data)
    leases = Lease.all_objects.filter(unit=unit).exclude(status=Lease.Status.DRAFT)
    number = data.get("lease_number", "")
    lease = (leases.filter(number__iexact=number) if number
             else leases.filter(status=Lease.Status.ACTIVE, archived_at__isnull=True)).first()
    if lease is None:
        message = (_("%(unit)s has no lease %(number)s.") if number
                   else _("%(unit)s has no active lease. Add lease_number for an ended one."))
        raise ValidationError(message % {"unit": unit.payment_reference, "number": number})
    if invoicing.opening_balance(lease) is not None:
        raise ValidationError(_("Lease %(number)s already has a balance brought forward. Change it on its account "
                                "page.") % {"number": lease.number})
    try:
        entry = invoicing.set_opening_balance(actor, lease, amount=data["amount"], as_of=_date(data["as_of"]),
                                              reason=data.get("reason", ""), request=request)
    except PermissionDenied:
        raise ValidationError(_("You cannot enter balances for %(name)s.") % {"name": unit.property.name}) from None
    if entry is None:
        raise ValidationError(_("The amount is zero. Leave the row out instead."))
    return entry, f"{unit.payment_reference} · {lease.number}"


def _whole(value: str, field: str, low: int, high: int) -> int:
    try:
        number = int(value)
    except ValueError:
        number = None
    if number is None or not low <= number <= high:
        raise ValidationError(_("%(field)s must be a whole number from %(low)s to %(high)s.")
                              % {"field": field, "low": low, "high": high})
    return number


def _create_lease(actor: Membership, data: dict, request, ctx):
    """A tenant (the one with this phone, or a new one) on a lease for the unit: a draft, or issued if asked."""
    unit = _unit(actor, data)
    start = _date(data["start_date"])
    rent, deposit = _money(data["rent"]), _money(data.get("deposit", ""))
    due_day = _whole(data["due_day"], "due_day", 1, 28) if data.get("due_day") else 1
    taken = Lease.all_objects.filter(unit=unit, status__in=[Lease.Status.DRAFT, Lease.Status.ACTIVE]).first()
    if taken is not None:
        raise ValidationError(_("%(unit)s already has a lease (%(lease)s).")
                              % {"unit": unit.payment_reference, "lease": taken.number or _("draft")})
    try:
        phone = normalize_phone(data["phone"])
    except InvalidPhoneNumber as exc:
        raise ValidationError(str(exc)) from None
    tenant = tenant_services.tenants_with_phone(actor.organization, phone).first()
    if tenant is None:
        tenant = tenant_services.create_tenant(actor, name=data["name"], phone=phone, email=data.get("email", ""),
                                               request=request)
        ctx["tenants"].append(tenant.pk)
    lease = lease_services.create_lease(actor, unit=unit, tenants=[tenant], start_date=start, rent=rent,
                                        deposit_amount=deposit or 0, due_day=due_day, request=request)
    if ctx["activate"]:
        lease = lease_services.activate_lease(actor, lease, request=request)
    return lease, f"{unit.payment_reference} · {tenant.name}"


CREATE = {Kind.UNITS: _create_unit, Kind.TENANTS: _create_tenant, Kind.BALANCES: _create_balance,
          Kind.LEASES: _create_lease}
MODEL = {Kind.UNITS: Unit, Kind.TENANTS: Tenant}


def _messages(exc: ValidationError) -> list[str]:
    if not hasattr(exc, "error_dict"):
        return exc.messages
    return [m if field == "__all__" else f"{field.replace('_', ' ')}: {m}"
            for field, msgs in exc.message_dict.items() for m in msgs]


def _process(actor: Membership, kind: str, rows: list[dict], request,
             activate: bool = False) -> tuple[list[dict], list[int], list[int]]:
    """Runs every row. Returns (rows with their errors, ids created, ids of tenants created along the way)."""
    checked, created, tenants = [], [], []
    for row in rows:
        errors, label = list(row.get("parse_errors", row["errors"])), ""
        ctx = {"activate": activate, "tenants": []}
        if not errors:
            try:
                with transaction.atomic():
                    obj, label = CREATE[kind](actor, row["data"], request, ctx)
            except ValidationError as exc:
                errors = _messages(exc)
            except PermissionDenied:
                errors = [_("You are not allowed to add this.")]
            else:
                created.append(obj.pk)
                tenants.extend(ctx["tenants"])
        checked.append({**row, "errors": errors, "label": label})
    return checked, created, tenants


def _counts(rows):
    ok = sum(1 for r in rows if not r["errors"])
    return ok, len(rows) - ok


# ---------------------------------------------------------------------------
# Actions
# ---------------------------------------------------------------------------


def preview_import(actor: Membership, kind: str, upload, request=None, activate: bool = False) -> ImportBatch:
    """Checks every row of a CSV or Excel file without keeping anything, and stores a PREVIEW batch."""
    require(actor, CAPABILITY[kind])
    rows = parse_upload(kind, upload)
    return preview_rows(actor, kind, rows, file_name=getattr(upload, "name", ""), activate=activate)


def preview_rows(actor: Membership, kind: str, rows: list[dict], *, file_name: str = "", activate: bool = False,
                 source: str = ImportBatch.Source.FILE) -> ImportBatch:
    """The same check for rows already read (from a file or a photo)."""
    require(actor, CAPABILITY[kind])
    activate = bool(activate and kind == Kind.LEASES and can(actor, "leases.activate"))
    for row in rows:
        row["parse_errors"] = row["errors"]
    with transaction.atomic():
        checked, _created, _tenants = _process(actor, kind, rows, None, activate)
        transaction.set_rollback(True)
    ok, bad = _counts(checked)
    return ImportBatch.objects.create(
        organization=actor.organization, kind=kind, created_by=actor.user, source=source, activate=activate,
        file_name=file_name[:200], rows=checked, ok_count=ok, error_count=bad,
    )


def _locked(actor: Membership, batch: ImportBatch) -> ImportBatch:
    if batch.organization_id != actor.organization_id:
        raise PermissionDenied(_("That record belongs to another organization."))
    require(actor, CAPABILITY[batch.kind])
    if not can_access(actor, batch):
        raise PermissionDenied(_("Only the person who uploaded this import can use it."))
    return ImportBatch.objects.select_for_update().get(pk=batch.pk)


def _scrub(rows):
    return [{**r, "data": {k: v for k, v in r["data"].items() if k not in SENSITIVE}} for r in rows]


@transaction.atomic
def apply_import(actor: Membership, batch: ImportBatch, request=None, *, checked_by_eye: bool = False) -> ImportBatch:
    locked = _locked(actor, batch)
    if locked.status != Status.PREVIEW:
        raise ValidationError(_("This import has already been applied or discarded."))
    if is_expired(locked):
        raise ValidationError(_("This preview is more than a day old. Upload the file again."))
    if locked.source == ImportBatch.Source.PHOTO and not checked_by_eye:
        raise ValidationError(_("Tick the box to confirm you checked the rows against the paper."))
    checked, created, tenants = _process(actor, locked.kind, locked.rows, request, locked.activate)
    ok, bad = _counts(checked)
    locked.rows, locked.created_ids, locked.ok_count, locked.error_count = _scrub(checked), created, ok, bad
    locked.created_tenant_ids = tenants
    locked.status, locked.applied_at = Status.APPLIED, timezone.now()
    locked.save()
    audit.record("import.apply", actor=actor.user, organization=actor.organization, obj=locked, request=request,
                 changes={"kind": [None, locked.kind], "created": [None, ok], "skipped": [None, bad]})
    return locked


@transaction.atomic
def discard_import(actor: Membership, batch: ImportBatch) -> ImportBatch:
    locked = _locked(actor, batch)
    if locked.status != Status.PREVIEW:
        raise ValidationError(_("Only an import that has not been applied can be discarded."))
    locked.status, locked.rows = Status.DISCARDED, []
    locked.save(update_fields=["status", "rows", "updated_at"])
    return locked


def _in_use(kind: str, obj) -> bool:
    return obj.leases.exists() if kind == Kind.UNITS else obj.lease_links.exists()


def _undo_balances(actor: Membership, batch: ImportBatch, request) -> tuple[list[str], list[str]]:
    """Reverses each imported balance that is still the lease's balance brought forward."""
    removed, kept = [], []
    entries = (LedgerEntry.objects.filter(organization=actor.organization, pk__in=batch.created_ids,
                                          kind=LedgerEntry.Kind.OPENING_BALANCE)
               .select_related("lease__unit__property"))
    for entry in entries:
        lease = entry.lease
        label = f"{lease.unit.payment_reference} · {lease.number}"
        if invoicing.opening_balance(lease) != entry:
            kept.append(label)  # changed by hand since the import
            continue
        try:
            with transaction.atomic():
                invoicing.set_opening_balance(actor, lease, amount=0, as_of=entry.entry_date, request=request)
        except (PermissionDenied, ValidationError):
            kept.append(label)
        else:
            removed.append(label)
    return removed, kept


def _undo_leases(actor: Membership, batch: ImportBatch, request) -> tuple[list[str], list[str]]:
    """Deletes the draft leases the batch made, then the tenants it made that no lease uses. Issued leases stay."""
    removed, kept = [], []
    leases = Lease.all_objects.filter(organization=actor.organization, pk__in=batch.created_ids)
    for lease in leases.select_related("unit__property"):
        label = f"{lease.unit.payment_reference} · {lease.number or _('draft')}"
        if lease.status != Lease.Status.DRAFT:
            kept.append(label)
            continue
        try:
            with transaction.atomic():
                lease_services.delete_draft_lease(actor, lease, request=request)
        except (PermissionDenied, ValidationError, ProtectedError):
            kept.append(label)
        else:
            removed.append(label)
    visible = set(tenant_services.visible_tenants(actor, Tenant.all_objects.filter(pk__in=batch.created_tenant_ids))
                  .values_list("pk", flat=True))
    for tenant in Tenant.all_objects.for_org(actor.organization).filter(pk__in=batch.created_tenant_ids):
        if tenant.pk not in visible or _in_use(Kind.TENANTS, tenant):
            kept.append(tenant.name)
            continue
        try:
            with transaction.atomic():
                tenant.delete()
        except ProtectedError:
            kept.append(tenant.name)
        else:
            removed.append(tenant.name)
    return removed, kept


@transaction.atomic
def undo_import(actor: Membership, batch: ImportBatch, request=None) -> tuple[list[str], list[str]]:
    """Deletes what the batch created, keeping anything already used. Returns (removed, kept) labels."""
    locked = _locked(actor, batch)
    if not locked.can_undo:
        raise ValidationError(_("An import can only be undone within 24 hours of applying it."))
    if locked.kind == Kind.BALANCES:
        removed, kept = _undo_balances(actor, locked, request)
        return _finish_undo(actor, locked, removed, kept, request)
    if locked.kind == Kind.LEASES:
        removed, kept = _undo_leases(actor, locked, request)
        return _finish_undo(actor, locked, removed, kept, request)
    model = MODEL[locked.kind]
    removed, kept = [], []
    visible = set(tenant_services.visible_tenants(actor, Tenant.all_objects.filter(pk__in=locked.created_ids))
                  .values_list("pk", flat=True)) if locked.kind == Kind.TENANTS else set()
    for obj in model.all_objects.for_org(actor.organization).filter(pk__in=locked.created_ids):
        label = obj.payment_reference if locked.kind == Kind.UNITS else obj.name
        if locked.kind == Kind.UNITS and not can(actor, "units.manage", obj.property):
            kept.append(label)
            continue
        if locked.kind == Kind.TENANTS and obj.pk not in visible:
            kept.append(label)
            continue
        if _in_use(locked.kind, obj):
            kept.append(label)
            continue
        try:
            with transaction.atomic():
                obj.delete()
        except ProtectedError:
            kept.append(label)
        else:
            removed.append(label)
    return _finish_undo(actor, locked, removed, kept, request)


def _finish_undo(actor, locked, removed, kept, request):
    locked.status, locked.undone_at = Status.UNDONE, timezone.now()
    locked.save(update_fields=["status", "undone_at", "updated_at"])
    audit.record("import.undo", actor=actor.user, organization=actor.organization, obj=locked, request=request,
                 changes={"removed": [None, removed], "kept": [None, kept]})
    return removed, kept


parse_csv = parse_upload  # the name the first version used; it reads Excel files too now
