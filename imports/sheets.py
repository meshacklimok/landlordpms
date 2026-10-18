"""Fill-in sheets for imports (D-156): an Excel workbook and a printable PDF, both pre-filled with what we already
know so the person only types what we do not.

- Tenants moving in: every vacant unit the member can see is listed with its asking rent. Leave a unit's tenant
  empty and the row is skipped on upload.
- Opening balances: every active lease is listed; fill in the amount owed.
- The printed sheet leaves out ID numbers and KRA PINs, because a photo of it goes to the model that reads it.
"""

import datetime
import io

from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

from accounts.models import Membership
from accounts.permissions import visible_properties
from core import brand
from leases.models import Lease
from payments.models import Payment
from properties.models import Property, Unit
from tenants.models import Tenant

from . import services
from .models import ImportBatch

Kind = ImportBatch.Kind
MAX_PREFILL = services.MAX_ROWS

# What the printed form asks for, in order. The keys are the import's own column names.
PRINT_COLUMNS = {
    Kind.UNITS: ("property_code", "unit_code", "building", "unit_type", "list_rent"),
    Kind.TENANTS: ("name", "phone", "alt_phone", "email", "emergency_contact_name", "emergency_contact_phone"),
    Kind.LEASES: ("property_code", "unit_code", "name", "phone", "rent", "deposit", "start_date"),
    Kind.BALANCES: ("property_code", "unit_code", "amount", "as_of"),
    Kind.PAYMENTS: ("property_code", "unit_code", "amount", "paid_at", "method", "reference"),
}
LABELS = {
    "property_code": _("Property code"), "unit_code": _("Unit"), "building": _("Building"), "unit_type": _("Type"),
    "list_rent": _("Rent"), "name": _("Tenant name"), "phone": _("Phone"), "alt_phone": _("Other phone"),
    "email": _("Email"), "emergency_contact_name": _("Emergency contact"),
    "emergency_contact_phone": _("Emergency phone"), "rent": _("Monthly rent"), "deposit": _("Deposit paid"),
    "start_date": _("Moved in (date)"), "amount": _("Owes now (KES)"), "as_of": _("As at (date)"),
    "paid_at": _("Date paid"), "method": _("Cash, M-Pesa, Bank"), "reference": _("M-Pesa code / reference"),
}
WIDTHS = {"property_code": 1.0, "unit_code": 1.0, "building": 1.3, "unit_type": 1.3, "list_rent": 1.1,
          "name": 2.6, "phone": 1.7, "alt_phone": 1.7, "email": 2.2, "emergency_contact_name": 2.2,
          "emergency_contact_phone": 1.7, "rent": 1.2, "deposit": 1.2, "start_date": 1.5, "amount": 1.4, "as_of": 1.5,
          "paid_at": 1.5, "method": 1.5, "reference": 2.0}
PAYMENT_LABELS = {"amount": _("Amount paid (KES)")}
TITLES = {
    Kind.UNITS: _("Units"), Kind.TENANTS: _("Tenants"), Kind.LEASES: _("Tenants moving in"),
    Kind.BALANCES: _("Opening balances"), Kind.PAYMENTS: _("Past payments"),
}
HINTS = {
    Kind.UNITS: _("One unit per line. Unit type: Apartment, House, Bedsitter, Studio, Shop, Office..."),
    Kind.TENANTS: _("One tenant per line. Phones like 0712 345 678."),
    Kind.LEASES: _("One tenant per line, on the line of the unit they live in. Leave a unit's line empty if it is "
                   "vacant. Dates like 31/01/2026."),
    Kind.BALANCES: _("What each tenant owed us when they started on NyumbaPMS. Put a minus sign if they paid "
                     "ahead."),
    Kind.PAYMENTS: _("One payment per line: what each tenant paid before you started on NyumbaPMS. Method is Cash, "
                     "M-Pesa, Bank or Cheque (Cash if left empty). Nobody is sent a message about these. If you "
                     "also enter an opening balance, only list payments made AFTER the balance date."),
}
CHOICES = {"unit_type": Unit.Type, "kind": Tenant.Kind, "id_type": Tenant.IdType, "method": Payment.Method}
TEXT_COLUMNS = {"property_code", "unit_code", "phone", "alt_phone", "emergency_contact_phone", "lease_number",
                "id_number", "kra_pin", "reference"}


def _units(actor: Membership, vacant_only: bool):
    units = (Unit.objects.filter(property__in=visible_properties(actor, Property.objects.all()))
             .select_related("property").order_by("property__code", "code"))
    if vacant_only:
        units = units.exclude(leases__status__in=[Lease.Status.DRAFT, Lease.Status.ACTIVE])
    return units[:MAX_PREFILL]


def prefill(actor: Membership, kind: str) -> list[dict]:
    """Rows we can fill in already: vacant units for move-ins, active leases for opening balances."""
    if kind == Kind.LEASES:
        return [{"property_code": u.property.code, "unit_code": u.code,
                 "rent": f"{u.list_rent:f}" if u.list_rent else ""} for u in _units(actor, vacant_only=True)]
    if kind in (Kind.BALANCES, Kind.PAYMENTS):
        leases = (Lease.objects.filter(status=Lease.Status.ACTIVE, unit__property__in=visible_properties(
            actor, Property.objects.all())).select_related("unit__property").order_by("unit__property__code",
                                                                                   "unit__code"))[:MAX_PREFILL]
        return [{"property_code": ln.unit.property.code, "unit_code": ln.unit.code, "lease_number": ln.number}
                for ln in leases]
    return []


def _instructions(kind: str) -> list[str]:
    columns = services.COLUMNS[kind]
    lines = [str(HINTS[kind]), "",
             _("Columns marked * are required: %(cols)s.") % {"cols": ", ".join(services.REQUIRED[kind])}, ""]
    lines += [f"{c}{'*' if c in services.REQUIRED[kind] else ''}: {services.EXAMPLES[kind][i] or '-'}"
              for i, c in enumerate(columns)]
    for column, choices in CHOICES.items():
        if column in columns:
            lines += ["", _("%(col)s can be: %(values)s") % {"col": column, "values": ", ".join(choices.values)}]
    return lines


def template_xlsx(actor: Membership, kind: str) -> bytes:
    """A workbook with the import's columns, dropdowns for the choices, phones kept as text, and a help sheet."""
    book = Workbook()
    sheet = book.active
    sheet.title = str(TITLES[kind])[:31]
    columns = services.COLUMNS[kind]
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="146C4A")
    for i, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=i, value=column)
        cell.font, cell.fill = head, fill
        dimension = sheet.column_dimensions[get_column_letter(i)]
        dimension.width = max(14, len(column) + 4)
        if column in TEXT_COLUMNS:
            dimension.number_format = "@"  # so Excel keeps the 0 of 0712...
        if column in CHOICES:
            options = ",".join(CHOICES[column].values)
            check = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
            sheet.add_data_validation(check)
            check.add(f"{get_column_letter(i)}2:{get_column_letter(i)}{MAX_PREFILL + 1}")
    sheet.freeze_panes = "A2"
    for r, values in enumerate(prefill(actor, kind), start=2):
        for i, column in enumerate(columns, start=1):
            if values.get(column):
                sheet.cell(row=r, column=i, value=values[column])
    help_sheet = book.create_sheet(str(_("How to fill in"))[:31])
    help_sheet.column_dimensions["A"].width = 110
    for r, line in enumerate(_instructions(kind), start=1):
        help_sheet.cell(row=r, column=1, value=line).alignment = Alignment(wrap_text=True)
    help_sheet["A1"].font = Font(bold=True)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _label(kind: str, column: str):
    return PAYMENT_LABELS.get(column, LABELS[column]) if kind == Kind.PAYMENTS else LABELS[column]


def template_pdf(actor: Membership, kind: str, today: datetime.date | None = None) -> bytes:
    """A printable form: the columns as headings, the rows we already know written in, blank lines under them."""
    today = today or timezone.localdate()
    columns = PRINT_COLUMNS[kind]
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"].clone("cell", fontSize=9, leading=11)
    head_style = styles["BodyText"].clone("head", fontSize=8, leading=10, fontName="Helvetica-Bold")
    page_width = landscape(A4)[0] - 24 * mm
    total = sum(WIDTHS[c] for c in columns)
    widths = [page_width * WIDTHS[c] / total for c in columns]
    known = prefill(actor, kind)
    body = [[Paragraph(str(r.get(c) or ""), cell_style) for c in columns] for r in known]
    body += [[""] * len(columns) for _line in range(5 if known else 20)]
    rows = [[Paragraph(str(_label(kind, c)), head_style) for c in columns]] + body
    table = Table(rows, colWidths=widths, rowHeights=[10 * mm] + [9 * mm] * len(body), repeatRows=1)
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, brand.MUTED),
        ("BACKGROUND", (0, 0), (-1, 0), brand.TINT),
        ("LINEBELOW", (0, 0), (-1, 0), 1.2, brand.PRIMARY),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    org_name = actor.organization.name
    doc_buffer = io.BytesIO()
    doc = brand.Doc(doc_buffer, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                    topMargin=12 * mm, bottomMargin=16 * mm, title=f"{org_name} · {TITLES[kind]}", author="NyumbaPMS")
    org = actor.organization
    intro = [
        brand.letterhead(org.display_name, org=org, kind=str(TITLES[kind]), number=f"{today:%d/%m/%Y}",
                         lines=brand.org_lines(org), width=landscape(A4)[0] - 24 * mm),
        Spacer(1, 4 * mm),
        Paragraph(str(HINTS[kind]), styles["BodyText"]),
        Paragraph(_("Write clearly in CAPITAL LETTERS and one line per entry. When it is filled in, take a photo of "
                    "each page (flat, close up, in good light) and upload it on NyumbaPMS under Import, or type it "
                    "in."),
                  styles["BodyText"]),
        Spacer(1, 4 * mm),
    ]
    doc.build(intro + [table])
    return doc_buffer.getvalue()
