"""Excel uploads, the move-in sheet (tenant + lease), the fill-in sheets and photos of a filled form (D-156)."""

import datetime
import io

import pytest
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from accounts.tests.factories import make_org, make_property
from accounts.tests.test_isolation import login
from ai_assistant.client import FakeClient
from imports import scan, services, sheets
from imports.models import ImportBatch
from leases.models import Lease
from meters.models import Scan
from meters.tests.conftest import jpeg
from properties import services as property_services
from tenants.models import Tenant

pytestmark = pytest.mark.django_db

LEASES, TENANTS, BALANCES = ImportBatch.Kind.LEASES, ImportBatch.Kind.TENANTS, ImportBatch.Kind.BALANCES
HEADER = ["property_code", "unit_code", "name", "phone", "rent", "start_date"]


@pytest.fixture
def owner():
    return make_org()


@pytest.fixture
def prop(owner):
    prop = make_property(owner.organization)
    for code, rent in (("A1", "15000"), ("A2", "16000"), ("A3", "")):
        property_services.create_unit(owner, prop, code=code, list_rent=rent or None)
    return prop


def xlsx(rows, name="tenants.xlsx", header=HEADER):
    book = Workbook()
    sheet = book.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return SimpleUploadedFile(name, out.getvalue())


def moving_in(prop, *rows):
    return xlsx([[prop.code, *r] for r in rows])


def errors(batch):
    return {r["line"]: r["errors"] for r in batch.rows if r["errors"]}


def test_an_excel_file_is_read_like_a_csv(owner, prop):
    # Excel hands over dates as dates, rent as numbers and drops the 0 of a phone typed as a number.
    upload = moving_in(prop, ["A1", "Wanjiku Kamau", 712345678, 15000.0, datetime.date(2026, 1, 1)])
    rows = services.parse_upload(LEASES, upload)
    assert rows[0]["data"] == {"property_code": prop.code, "unit_code": "A1", "name": "Wanjiku Kamau",
                               "phone": "0712345678", "rent": "15000", "start_date": "2026-01-01"}
    assert rows[0]["line"] == 2 and not rows[0]["errors"]


def test_a_broken_excel_file_is_refused_kindly(owner):
    with pytest.raises(ValidationError, match="Excel"):
        services.parse_upload(LEASES, SimpleUploadedFile("x.xlsx", b"not a zip"))
    with pytest.raises(ValidationError, match=r"\.xls files"):
        services.parse_upload(LEASES, SimpleUploadedFile("x.xls", b"old"))


def test_moving_in_checks_everything_and_changes_nothing(owner, prop):
    batch = services.preview_import(owner, LEASES, moving_in(
        prop,
        ["A1", "Wanjiku Kamau", "0712345678", 15000, "01/01/2026"],
        ["A1", "Someone Else", "0722000000", 15000, "01/01/2026"],  # A1 twice in one file
        ["A2", "Otieno", "12345", 16000, "2026-01-01"],              # not a phone
        ["A9", "Nobody", "0733000000", 1000, "2026-01-01"],          # no such unit
    ))
    assert batch.ok_count == 1 and batch.error_count == 3
    assert "already has a lease" in errors(batch)[3][0]
    assert 4 in errors(batch) and "No unit A9" in errors(batch)[5][0]
    assert not Lease.all_objects.exists() and not Tenant.all_objects.exists()


def test_moving_in_makes_tenants_and_draft_leases_and_undo_removes_them(owner, prop):
    existing = Tenant.all_objects.create(organization=owner.organization, name="Old", phone="+254722000000")
    batch = services.preview_import(owner, LEASES, moving_in(
        prop, ["A1", "Wanjiku Kamau", "0712345678", 15000, "2026-01-01"],
        ["A2", "Same Person", "0722000000", 16000, "2026-01-01"]))
    batch = services.apply_import(owner, batch)
    leases = Lease.all_objects.order_by("unit__code")
    assert [(lease.unit.code, lease.status) for lease in leases] == [("A1", "DRAFT"), ("A2", "DRAFT")]
    assert leases[1].lease_tenants.get().tenant == existing  # the phone matched, so no second tenant
    assert Tenant.all_objects.count() == 2 and len(batch.created_tenant_ids) == 1
    removed, kept = services.undo_import(owner, batch)
    assert len(removed) == 3 and not kept
    assert not Lease.all_objects.exists() and list(Tenant.all_objects.all()) == [existing]


def test_issuing_the_leases_keeps_them_when_undone(owner, prop):
    batch = services.preview_import(
        owner, LEASES, moving_in(prop, ["A1", "Wanjiku Kamau", "0712345678", 15000, "2026-01-01"]), activate=True)
    assert batch.activate
    batch = services.apply_import(owner, batch)
    lease = Lease.all_objects.get()
    assert lease.status == "ACTIVE" and lease.number.startswith("LSE")
    removed, kept = services.undo_import(owner, batch)
    assert not removed and kept and Lease.all_objects.filter(pk=lease.pk).exists()


def test_the_excel_sheet_lists_vacant_units_and_blank_rows_are_skipped(owner, prop):
    first = moving_in(prop, ["A1", "Wanjiku", "0712345678", 15000, "2026-01-01"])
    services.apply_import(owner, services.preview_import(owner, LEASES, first))
    book = load_workbook(io.BytesIO(sheets.template_xlsx(owner, LEASES)))
    sheet = book.worksheets[0]
    rows = [[c.value for c in r] for r in sheet.iter_rows(min_row=2, max_col=2) if r[1].value]
    assert [r[1] for r in rows] == ["A2", "A3"]  # A1 has a lease already
    assert sheet.cell(row=2, column=6).value == "16000.00" and len(book.worksheets) == 2
    # Fill in one tenant and upload the same workbook: the empty vacant unit is ignored.
    sheet.cell(row=2, column=3, value="Otieno")
    sheet.cell(row=2, column=4, value="0722000000")
    sheet.cell(row=2, column=8, value="2026-02-01")
    out = io.BytesIO()
    book.save(out)
    parsed = services.parse_upload(LEASES, SimpleUploadedFile("sheet.xlsx", out.getvalue()))
    assert [r["data"]["unit_code"] for r in parsed] == ["A2"] and not parsed[0]["errors"]


def test_balances_sheet_lists_active_leases(owner, prop):
    batch = services.preview_import(
        owner, LEASES, moving_in(prop, ["A1", "Wanjiku", "0712345678", 15000, "2026-01-01"]), activate=True)
    services.apply_import(owner, batch)
    sheet = load_workbook(io.BytesIO(sheets.template_xlsx(owner, BALANCES))).worksheets[0]
    assert [r[1].value for r in sheet.iter_rows(min_row=2, max_col=3)] == ["A1"]


def test_the_printable_form_is_a_pdf_without_id_columns(owner, prop):
    for kind in sheets.PRINT_COLUMNS:
        assert sheets.template_pdf(owner, kind).startswith(b"%PDF")
    assert not {"id_number", "kra_pin"} & {c for cols in sheets.PRINT_COLUMNS.values() for c in cols}


def filled_form(*rows):
    columns = sheets.PRINT_COLUMNS[LEASES]
    return FakeClient.image({"legible": True, "rows": [dict(zip(columns, r, strict=True)) for r in rows]})


def test_a_photo_of_the_form_becomes_a_preview_that_must_be_checked_by_eye(owner, prop):
    FakeClient.script = [filled_form(
        (prop.code, "A1", "WANJIKU KAMAU", "0712345678", "15000", None, "01/01/2026"),
        (prop.code, "A2", None, None, "16000", None, None),   # a vacant unit nobody wrote on
    )]
    rows = scan.read_form(owner, LEASES, jpeg())
    assert len(rows) == 1 and rows[0]["data"]["name"] == "WANJIKU KAMAU"
    call = FakeClient.calls[0]
    assert "property_code" in call["instructions"] and call["schema"]["properties"]["rows"]["type"] == "array"
    assert Scan.objects.get().kind == "SHEET"
    batch = services.preview_rows(owner, LEASES, rows, source=ImportBatch.Source.PHOTO)
    assert batch.ok_count == 1 and not Lease.all_objects.exists()
    with pytest.raises(ValidationError, match="checked"):
        services.apply_import(owner, batch)
    services.apply_import(owner, batch, checked_by_eye=True)
    assert Lease.all_objects.count() == 1


def test_a_photo_with_nothing_readable_is_refused(owner, prop):
    FakeClient.script = [FakeClient.image({"legible": False, "rows": []})]
    with pytest.raises(ValidationError, match="No lines"):
        scan.read_form(owner, LEASES, jpeg())
    assert Scan.objects.get().status == "UNCLEAR"


def test_pages(client, owner, prop):
    login(client, owner)
    assert client.get(reverse("imports:list")).status_code == 200
    r = client.get(reverse("imports:template_xlsx", args=["leases"]))
    assert r.status_code == 200 and r.content[:2] == b"PK"
    r = client.get(reverse("imports:print", args=["tenants"]))
    assert r["Content-Type"] == "application/pdf" and r.content.startswith(b"%PDF")
    r = client.post(reverse("imports:upload", args=["leases"]),
                    {"file": moving_in(prop, ["A1", "Wanjiku", "0712345678", 15000, "2026-01-01"]), "activate": "1"})
    detail = r.url
    assert client.get(detail).status_code == 200
    FakeClient.script = [filled_form((prop.code, "A2", "OTIENO", "0722000000", "16000", None, "2026-01-01"))]
    r = client.post(reverse("imports:photo", args=["leases"]), {"photo": jpeg()})
    page = client.get(r.url)
    assert b"checked these rows against the paper" in page.content
    r = client.post(r.url, {"action": "apply"})
    assert not Lease.all_objects.exists()  # not ticked, so nothing was added
